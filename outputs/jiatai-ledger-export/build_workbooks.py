import csv
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path("/Users/aslight/Documents/workspace/IDEAWorkspace/workhub-server/outputs/jiatai-ledger-export")
OUT = BASE / "final"
OUT.mkdir(parents=True, exist_ok=True)

JOBS = [
    (
        BASE / "loan-ledger.csv",
        "放款台账",
        OUT / "嘉泰资产平台_全量放款台账_20260529.xlsx",
    ),
    (
        BASE / "prepayment-ledger.csv",
        "提前结清台账",
        OUT / "嘉泰资产平台_提前结清台账_结算日期20260427-20260529.xlsx",
    ),
]

DATE_HEADERS = {"放款时间", "起租日", "到期日", "起息日", "最晚还款日", "宽限期到期日", "结算时间", "用户实际还款时间"}
TEXT_HEADERS = {"客户身份证号", "放款收款账户账号", "联行号", "用户编号", "放款申请单号", "用户账单编号", "进件申请单号", "代扣渠道流水", "产品编号", "合作方编号"}


def coerce(value: str, header: str):
    if value == "":
        return None
    if header in TEXT_HEADERS:
        return value
    if header in DATE_HEADERS:
        for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(value, fmt)
                return parsed.date() if fmt == "%Y-%m-%d" else parsed
            except ValueError:
                pass
        return value
    numeric_hint = any(token in header for token in ("金额", "本金", "利息", "手续费", "服务费", "担保费", "罚息", "期数", "天数"))
    if numeric_hint:
        try:
            number = Decimal(value)
            return int(number) if number == number.to_integral_value() else float(number)
        except (InvalidOperation, ValueError):
            return value
    return value


def build(csv_path: Path, sheet_name: str, xlsx_path: Path):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_name)
    ws.freeze_panes = "A2"

    with csv_path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        headers = next(reader)
        header_cells = []
        fill = PatternFill("solid", fgColor="D9EAF7")
        thin = Side(style="thin", color="D9D9D9")
        border = Border(bottom=thin)
        for header in headers:
            cell = WriteOnlyCell(ws, value=header)
            cell.font = Font(name="Microsoft YaHei", bold=True, size=10)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            header_cells.append(cell)
        ws.append(header_cells)
        row_count = 0
        for row in reader:
            ws.append([coerce(value, headers[idx]) for idx, value in enumerate(row)])
            row_count += 1

    col_count = len(headers)
    ws.auto_filter.ref = f"A1:{get_column_letter(col_count)}{row_count + 1}"

    for idx, header in enumerate(headers, start=1):
        width = 14
        if len(header) >= 8:
            width = 18
        if header in TEXT_HEADERS or header in DATE_HEADERS:
            width = max(width, 20)
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 32)

    wb.save(xlsx_path)

    check = load_workbook(xlsx_path, read_only=True, data_only=True)
    check_ws = check[sheet_name]
    verified_total_rows = 0
    verified_cols = 0
    for verified_total_rows, row in enumerate(check_ws.iter_rows(values_only=True), start=1):
        if verified_total_rows == 1:
            verified_cols = len(row)
    verified_rows = max(verified_total_rows - 1, 0)
    check.close()
    print(f"{xlsx_path.name}\trows={row_count}\tcols={col_count}\tverified_rows={verified_rows}\tverified_cols={verified_cols}")


for job in JOBS:
    build(*job)
