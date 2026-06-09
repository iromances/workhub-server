from __future__ import annotations
from openpyxl import load_workbook
from pathlib import Path
import re, json, hashlib, time, sys
from collections import Counter
from urllib import request, parse, error

BASE='http://127.0.0.1:8081/api'
XLSX=Path('/Users/aslight/Library/CloudStorage/OneDrive-个人/公司/支付渠道/三方渠道生产环境.xlsx')
INCLUDE=['资产平台','保费分期','设备','交通银行','账单管理','嘉泰']
EXCLUDE=['微信小程序、公众号','通联全民测试商户号']
BUS_CODES={
    '资产平台':'ASSET_PLATFORM',
    '保费分期':'PREMIUM_INSTALLMENT',
    '设备':'EQUIPMENT',
    '嘉泰汇浦设备':'JIATAI_HUIPU_EQUIPMENT',
    '嘉泰资产平台':'JIATAI_ASSET_PLATFORM',
    '账单管理':'BILLING_MANAGEMENT',
    '可信资产交易平台':'TRUSTED_ASSET_TRADING',
    '可信--租赁保理':'TRUSTED_LEASING_FACTORING',
    '可信--卡尔':'TRUSTED_CARL',
    '可信--账单管理':'TRUSTED_BILLING',
}
SHEET_DEFAULT_BIZ={'设备':'设备','账单管理':'账单管理'}
CHANNELS={
    '宝付':('BAOFOO','宝付','宝付'),
    '易宝':('YEEPAY','易宝','易宝'),
    '交通银行':('BOCOM','交通银行','交通银行'),
    '企业微信':('WECOM_PAY','企业微信','企业微信'),
    '企业支付宝':('ALIPAY_ENTERPRISE','企业支付宝','支付宝'),
    '通联（通联通）':('ALLINPAY_TLT','通联通','通联'),
    '通联（收银宝）':('ALLINPAY_SYB','通联收银宝','通联'),
    '通联通':('ALLINPAY_TLT','通联通','通联'),
    '通联':('ALLINPAY','通联','通联'),
}
PURPOSES={
    '代收':['WITHHOLD'],'代付':['PAY_OUT'],'签约':['SIGN_AGREEMENT'],'绑卡':['BIND_CARD'],'分账':['SPLIT_SETTLEMENT'],
    '代收分账':['WITHHOLD_SPLIT_SETTLEMENT'],'退款主户':['REFUND_MAIN_ACCOUNT'],
    '代收子商户（生产测试用）':['WITHHOLD_SUB_MERCHANT_PROD_TEST'],'代收子商户（北京）':['WITHHOLD_SUB_MERCHANT_BEIJING'],'代收子商户（天津）':['WITHHOLD_SUB_MERCHANT_TIANJIN'],'代收子商户':['WITHHOLD_SUB_MERCHANT'],
    '先行通被分账商户':['SPLIT_RECEIVER_XXT'],'里易被分账商户':['SPLIT_RECEIVER_LIYI'],'账户体系':['ACCOUNT_SYSTEM'],'签约&代收':['SIGN_AGREEMENT','WITHHOLD'],
}
RELATION_COLS={'共享协议商户':'SHARED_AGREEMENT','正常分账商户':'NORMAL_SPLIT','代偿回购':'COMPENSATION_REPURCHASE','分账商户':'SPLIT_RECEIVER'}
SENSITIVE_MAP={'私钥密码':'PRIVATE_KEY_PASSWORD','私钥 / 密码':'PRIVATE_KEY_PASSWORD','用户名（接口用）':'USERNAME','用户名':'USERNAME','用户密码（接口用）':'PASSWORD','用户密码':'PASSWORD','AESKEY':'AES_KEY','密钥':'SIGN_KEY','账户名':'USERNAME','账户密码':'PASSWORD','交易密码':'TRADE_PASSWORD','pin码':'PASSWORD','安全密码':'PASSWORD'}
PARAM_NAMES={'手续费','APPID','业务代码','终端号','备注','密钥类型','企业代码','用户号','ip'}

def clean(v):
    if v is None: return ''
    s=str(v).replace('\xa0',' ').strip()
    if s.endswith('.0') and re.fullmatch(r'\d+\.0', s): s=s[:-2]
    return s.strip()

def normalize_business(biz, sheet=None, remark=None):
    biz=clean(biz)
    if biz in {'叉车','智能柜'}: return '设备'
    if biz=='汇浦设备金融': return '嘉泰汇浦设备'
    if biz=='趣学呗': return '账单管理'
    if not biz and sheet in SHEET_DEFAULT_BIZ: return SHEET_DEFAULT_BIZ[sheet]
    return biz

def code_from_name(name):
    if name in BUS_CODES: return BUS_CODES[name]
    return 'BIZ_'+hashlib.md5(name.encode()).hexdigest()[:8].upper()

def parse_project(text,biz):
    text=clean(text)
    if not text: return None
    if '/' in text:
        code,name=text.split('/',1); return clean(code), clean(name)
    slug=re.sub(r'[^A-Za-z0-9]+','_',text).strip('_').upper()
    if slug and re.search(r'[A-Za-z]', slug): return slug, text
    return code_from_name(biz)+'_'+hashlib.md5(text.encode()).hexdigest()[:6].upper(), text

def parse_merchant(text):
    text=clean(text)
    if not text or text in {'同上','复用强云','未配置'}: return None
    m=re.search(r'([A-Za-z0-9][A-Za-z0-9_\-]{3,})', text)
    if not m: return None
    code=m.group(1).strip(); name=text[m.end():].strip(' /\t-') or code
    return code,name,text

def extract_merchants_from_text(text):
    text=clean(text)
    if not text: return []
    parts=re.split(r'\s*/\s*|\n|；|;', text)
    out=[]
    for p in parts:
        pm=parse_merchant(p)
        if pm: out.append(pm)
    return out

def normalize_channel(raw,last=''):
    raw=clean(raw)
    if raw=='同上': raw=last
    if not raw: raw=last
    return CHANNELS.get(raw, (re.sub(r'\W+','_',raw).strip('_').upper() or 'UNKNOWN', raw or '未知渠道', raw or '未知'))

def normalize_purpose(raw):
    raw=clean(raw)
    if raw in PURPOSES: return PURPOSES[raw]
    if '代付' in raw and '收款' in raw: return ['PAY_OUT','WITHHOLD']
    return []

def header_map(ws,rownum):
    rows=list(ws.iter_rows(min_row=rownum,max_row=rownum+1,values_only=True))
    h=[clean(v) for v in rows[0]]; sub=[clean(v) for v in rows[1]] if len(rows)>1 else []
    mp={}
    for i,x in enumerate(h):
        name=x or (sub[i] if i<len(sub) else '')
        if name in {'','渠道公钥','商户私钥'}: continue
        if name not in mp: mp[name]=i
    return mp

def rowval(row,mp,key):
    i=mp.get(key); return clean(row[i]) if i is not None and i<len(row) else ''

def parse_data():
    wb=load_workbook(XLSX,data_only=True,read_only=True)
    groups={}; projects={}; channels={}; merchants={}; merchant_remarks={}; credentials=[]; params=[]; binding_map={}; relation_raw=[]; skipped=[]
    for s in INCLUDE:
        ws=wb[s]
        if s=='交通银行':
            mp=header_map(ws,1)
            for rn,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
                acct=re.sub(r'\s+','',rowval(row,mp,'账户号') or rowval(row,mp,'pin码'))
                name=rowval(row,mp,'账户名')
                if not acct and not name: continue
                remark=rowval(row,mp,'备注')
                biz='账单管理' if remark=='趣学呗' else ('保费分期' if '保费分期' in remark else '资产平台')
                p=parse_project(remark or '交通银行账户',biz) or (code_from_name(biz)+'_BOCOM', remark or '交通银行账户')
                groups[biz]=code_from_name(biz); projects[p[0]]=(p[1],biz)
                ch=CHANNELS['交通银行']; channels[ch[0]]=ch
                merchants[(ch[0],acct)]=(name or acct,'PROD',None,'ACTIVE')
                merchant_remarks[(ch[0],acct)]=remark
                binding_map.setdefault((p[0],ch[0],acct),set()).update(['PAY_OUT','WITHHOLD'])
                for key in ['pin码','安全密码']:
                    val=rowval(row,mp,key)
                    if val: credentials.append((ch[0],acct,key,SENSITIVE_MAP[key],val,s,rn))
                for key in ['企业代码','用户号','ip','备注']:
                    val=rowval(row,mp,key)
                    if val: params.append((ch[0],acct,key,val,s,rn))
            continue
        header=3 if s=='嘉泰' else 1
        mp=header_map(ws,header); last_biz=''; last_project=''; last_channel=''; last_purpose=''
        for rn,row in enumerate(ws.iter_rows(min_row=header+1,values_only=True),header+1):
            if not any(clean(x) for x in row): continue
            biz=rowval(row,mp,'业务') or rowval(row,mp,'业务线')
            project_text=rowval(row,mp,'项目'); purpose=rowval(row,mp,'用途'); channel_raw=rowval(row,mp,'渠道'); merchant_raw=rowval(row,mp,'渠道商户')
            if s=='资产平台' and not purpose and project_text in PURPOSES:
                vals=[clean(x) for x in row]; biz=vals[1]; project_text=vals[2]; purpose=vals[3]; channel_raw=vals[4]; merchant_raw=vals[5]
            if biz: last_biz=normalize_business(biz,s)
            elif not last_biz and s in SHEET_DEFAULT_BIZ: last_biz=SHEET_DEFAULT_BIZ[s]
            biz=normalize_business(last_biz,s)
            if project_text: last_project=project_text
            project_text=project_text or last_project
            if purpose: last_purpose=purpose
            purpose=purpose or last_purpose
            ch=normalize_channel(channel_raw,last_channel)
            if ch[1] != '未知渠道' and channel_raw and channel_raw!='同上': last_channel=channel_raw
            merchant=parse_merchant(merchant_raw)
            if not merchant and not any(rowval(row,mp,k) for k in RELATION_COLS):
                if any([merchant_raw,purpose,project_text]): skipped.append({'sheet':s,'row':rn,'reason':'缺少业务/项目/用途/商户之一','business':biz,'project':project_text,'purpose':purpose,'merchant':merchant_raw})
                continue
            if not merchant: continue
            pcodes=normalize_purpose(purpose)
            if not pcodes:
                skipped.append({'sheet':s,'row':rn,'reason':'用途无法映射','purpose':purpose}); continue
            p=parse_project(project_text,biz)
            if not p:
                skipped.append({'sheet':s,'row':rn,'reason':'项目无法解析','project':project_text}); continue
            pcode,pname=p; mcode,mname,raw=merchant
            groups[biz]=code_from_name(biz); projects[pcode]=(pname,biz); channels[ch[0]]=ch
            appid=rowval(row,mp,'APPID') or None
            rem=rowval(row,mp,'备注')
            merchants[(ch[0],mcode)]=(mname,'PROD',appid,'ACTIVE')
            if rem: merchant_remarks[(ch[0],mcode)]=rem
            binding_map.setdefault((pcode,ch[0],mcode),set()).update(pcodes)
            for key in PARAM_NAMES:
                val=rowval(row,mp,key)
                if val: params.append((ch[0],mcode,key,val,s,rn))
            for key,typ in SENSITIVE_MAP.items():
                val=rowval(row,mp,key)
                if val: credentials.append((ch[0],mcode,key,typ,val,s,rn))
            for key,role in RELATION_COLS.items():
                val=rowval(row,mp,key)
                for rm in extract_merchants_from_text(val):
                    relation_raw.append((pcode,ch[0],mcode,rm[0],rm[1],role,key,s,rn))
                    merchants.setdefault((ch[0],rm[0]),(rm[1],'PROD',None,'ACTIVE'))
    return groups,projects,channels,merchants,merchant_remarks,credentials,params,binding_map,relation_raw,skipped

class Api:
    def __init__(self): self.token=None
    def call(self,method,path,data=None,params=None):
        url=BASE+path
        if params: url+='?'+parse.urlencode({k:v for k,v in params.items() if v is not None})
        body=None if data is None else json.dumps(data,ensure_ascii=False).encode('utf-8')
        req=request.Request(url, data=body, method=method, headers={'Content-Type':'application/json'})
        if self.token: req.add_header('Authorization','Bearer '+self.token)
        try:
            with request.urlopen(req,timeout=30) as resp:
                txt=resp.read().decode('utf-8')
        except error.HTTPError as e:
            txt=e.read().decode('utf-8')
            raise RuntimeError(f'{method} {path} HTTP {e.code}: {txt[:500]}')
        res=json.loads(txt) if txt else None
        if res and res.get('code')!=0:
            raise RuntimeError(f'{method} {path} API: {res}')
        return res.get('data') if res else None
    def login(self):
        data=self.call('POST','/auth/login',{'username':'admin','password':'admin123'})
        self.token=data['token']


def find_by(items, **kw):
    for x in items:
        if all(x.get(k)==v for k,v in kw.items()): return x
    return None

def main():
    groups,projects,channels,merchants,merchant_remarks,credentials,params,binding_map,relation_raw,skipped=parse_data()
    api=Api(); api.login()
    stats=Counter(); errors=[]
    # channels
    channel_ids={}
    current_channels=api.call('GET','/payment/channels',params={})['items']
    for code,(ccode,cname,vendor) in channels.items():
        found=find_by(current_channels, code=ccode)
        payload={'code':ccode,'name':cname,'vendorName':vendor,'status':'ACTIVE','description':'Excel 三方渠道生产环境导入'}
        try:
            if found:
                detail=api.call('PUT',f"/payment/channels/{found['id']}",payload); stats['channelsUpdated']+=1
            else:
                detail=api.call('POST','/payment/channels',payload); stats['channelsCreated']+=1
            channel_ids[ccode]=detail['id']
        except Exception as e: errors.append({'type':'channel','code':ccode,'error':str(e)})
    # projects
    project_ids={}
    current_projects=api.call('GET','/projects',params={})['items']
    for pcode,(pname,biz) in projects.items():
        found=find_by(current_projects, code=pcode)
        payload={'businessLineCode':code_from_name(biz),'businessLineName':biz,'code':pcode,'name':pname,'type':'业务项目','group':biz,'ownerUserName':'admin','status':'进行中','description':'Excel 三方渠道生产环境导入'}
        try:
            if found:
                detail=api.call('PUT',f"/projects/{found['id']}",payload); stats['projectsUpdated']+=1
            else:
                detail=api.call('POST','/projects',payload); stats['projectsCreated']+=1
            project_ids[pcode]=detail['id']
        except Exception as e: errors.append({'type':'project','code':pcode,'error':str(e)})
    # merchants
    merchant_ids={}
    current_merchants=api.call('GET','/payment/merchants',params={})['items']
    for (ccode,mcode),(mname,env,appid,status) in merchants.items():
        cid=channel_ids.get(ccode)
        if not cid: continue
        found=None
        for x in current_merchants:
            if x.get('channelId')==cid and x.get('merchantCode')==mcode and x.get('environment')==env:
                found=x; break
        payload={'channelId':cid,'merchantCode':mcode,'merchantName':mname,'environment':env,'status':status,'appId':appid,'settlementSubject':mname,'remark':merchant_remarks.get((ccode,mcode),'Excel 三方渠道生产环境导入')}
        try:
            if found:
                detail=api.call('PUT',f"/payment/merchants/{found['id']}",payload); stats['merchantsUpdated']+=1
            else:
                detail=api.call('POST','/payment/merchants',payload); stats['merchantsCreated']+=1
            merchant_ids[(ccode,mcode)]=detail['id']
        except Exception as e: errors.append({'type':'merchant','channel':ccode,'code':mcode,'error':str(e)})
    # params
    seen_params=set()
    for ccode,mcode,key,val,s,rn in params:
        mid=merchant_ids.get((ccode,mcode));
        if not mid or not val: continue
        pkey=re.sub(r'[^A-Za-z0-9_\-]+','_',key).strip('_') or key
        dedup=(mid,pkey,val)
        if dedup in seen_params: continue
        seen_params.add(dedup)
        try:
            api.call('POST',f'/payment/merchants/{mid}/params',{'paramKey':pkey,'paramValue':str(val),'valueType':'TEXT','sensitive':False,'remark':f'{s} 第{rn}行 {key}'})
            stats['paramsSaved']+=1
        except Exception as e: errors.append({'type':'param','merchant':mcode,'key':key,'error':str(e)})
    # credentials plain storage
    seen_cred=set()
    for ccode,mcode,key,typ,val,s,rn in credentials:
        mid=merchant_ids.get((ccode,mcode));
        if not mid or not val: continue
        ckey=re.sub(r'[^A-Za-z0-9_\-]+','_',key).strip('_') or key
        # include type if repeated under same key via different rows; API overwrites intentionally by key
        dedup=(mid,ckey,str(val))
        if dedup in seen_cred: continue
        seen_cred.add(dedup)
        try:
            api.call('POST',f'/payment/merchants/{mid}/credentials',{'credentialKey':ckey,'credentialName':key,'credentialType':typ,'credentialValue':str(val),'status':'ACTIVE','plainStorage':True,'remark':f'{s} 第{rn}行 {key}; 明文暂存，待后续重新加密'})
            stats['credentialsSaved']+=1
        except Exception as e: errors.append({'type':'credential','merchant':mcode,'key':key,'error':str(e)})
    # bindings with relations
    current_bindings=api.call('GET','/payment/bindings',params={})['items']
    relation_by_main={}
    for pcode,ccode,mcode,rmcode,rmname,role,key,s,rn in relation_raw:
        relation_by_main.setdefault((pcode,ccode,mcode),[]).append((ccode,rmcode,role,key,s,rn))
    for (pcode,ccode,mcode),pcodes in binding_map.items():
        pid=project_ids.get(pcode); mid=merchant_ids.get((ccode,mcode))
        if not pid or not mid: continue
        relations=[]
        for rcc,rmcode,role,key,s,rn in relation_by_main.get((pcode,ccode,mcode),[]):
            rmid=merchant_ids.get((rcc,rmcode))
            if rmid: relations.append({'merchantId':rmid,'relationRole':role,'relationName':key,'priority':len(relations)+1,'remark':f'{s} 第{rn}行'})
        # find existing by project + merchant; update broadest purposes
        found=None
        for b in current_bindings:
            if b.get('projectId')==pid and b.get('merchantId')==mid:
                found=b; break
        payload={'projectId':pid,'merchantId':mid,'purposeCode':sorted(pcodes)[0],'purposeCodes':sorted(pcodes),'priority':1,'defaultBinding':True,'status':'ACTIVE','remark':'Excel 三方渠道生产环境导入','relations':relations}
        try:
            if found:
                api.call('PUT',f"/payment/bindings/{found['id']}",payload); stats['bindingsUpdated']+=1
            else:
                api.call('POST','/payment/bindings',payload); stats['bindingsCreated']+=1
        except Exception as e: errors.append({'type':'binding','project':pcode,'merchant':mcode,'purposes':sorted(pcodes),'error':str(e)})
    summary={'stats':dict(stats),'skipped':skipped,'errors':errors[:200], 'errorCount':len(errors), 'parsed':{'businessLines':len(groups),'projects':len(projects),'channels':len(channels),'merchants':len(merchants),'bindings':len(binding_map),'credentials':len(credentials),'params':len(params),'relations':len(relation_raw)}}
    Path('outputs/payment-import/import-result.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    if errors: sys.exit(2)

if __name__=='__main__': main()
